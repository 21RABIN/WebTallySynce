from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path("/home/user/Videos/TALLY/Tally_BackUp/Correct Working Tally Synce/tallyconnector_Jun13/tallyconnector")
SOURCE_PPT = Path("/home/user/Downloads/Source_Code_and_Database_Version_Management_Presentation.pptx")
OUTPUT_PPT = Path("/home/user/Downloads/Source_Code_and_Database_Version_Management_Presentation_21RABIN_realtime.pptx")
GENERIC = ROOT / "output" / "ppt_assets"
REAL = ROOT / "output" / "github_real"


def clear_non_title_shapes(slide):
    title = slide.shapes.title
    for shape in list(slide.shapes):
        if shape == title:
            continue
        shape._element.getparent().remove(shape._element)


def add_textbox(slide, left, top, width, height, title, lines, title_size=26, body_size=18):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(title_size)
    p.font.bold = True
    p.font.color.rgb = RGBColor(31, 41, 55)
    p.alignment = PP_ALIGN.LEFT

    for line in lines:
        p = tf.add_paragraph()
        p.text = line
        p.level = 0
        p.font.size = Pt(body_size)
        p.font.color.rgb = RGBColor(55, 65, 81)
        p.space_after = Pt(8)


def add_picture(slide, path, left, top, width=None, height=None):
    slide.shapes.add_picture(str(path), left, top, width=width, height=height)


def main():
    prs = Presentation(str(SOURCE_PPT))

    # Slide 2
    slide = prs.slides[1]
    clear_non_title_shapes(slide)
    add_textbox(
        slide,
        Inches(0.55),
        Inches(1.45),
        Inches(4.15),
        Inches(5.3),
        "Real GitHub profile example",
        [
            "Public profile used: `github.com/21RABIN`.",
            "This shows repositories, stars, and contribution activity.",
            "Teams can review owner profile and public code work quickly.",
            "This makes the presentation use your real GitHub account view.",
        ],
    )
    add_picture(slide, REAL / "21RABIN_profile.png", Inches(4.85), Inches(1.4), width=Inches(7.35))

    # Slide 3
    slide = prs.slides[2]
    clear_non_title_shapes(slide)
    add_textbox(
        slide,
        Inches(0.5),
        Inches(1.42),
        Inches(4.1),
        Inches(5.25),
        "Real repo workflow example",
        [
            "Repository used: `21RABIN/billing`.",
            "Clone the repo, create branch, commit, and push.",
            "GitHub then shows the repo page and latest commit status.",
            "This is your real public repository screen.",
        ],
    )
    add_picture(slide, REAL / "21RABIN_billing_repo.png", Inches(4.82), Inches(1.35), width=Inches(7.25))

    # Slide 4
    slide = prs.slides[3]
    clear_non_title_shapes(slide)
    add_textbox(
        slide,
        Inches(0.5),
        Inches(1.42),
        Inches(4.1),
        Inches(5.25),
        "Branch view from your GitHub",
        [
            "GitHub branch screen for `21RABIN/billing`.",
            "The repository page also shows `3 Branches` near the branch selector.",
            "Developers work in branches, then merge back to `main`.",
            "This is a real branch-management example from your account.",
        ],
    )
    add_picture(slide, REAL / "21RABIN_billing_branches.png", Inches(4.9), Inches(1.55), width=Inches(7.0))
    add_picture(slide, GENERIC / "git_branch_graph_terminal.png", Inches(5.35), Inches(4.45), width=Inches(6.3))

    # Slide 5
    slide = prs.slides[4]
    clear_non_title_shapes(slide)
    add_textbox(
        slide,
        Inches(0.5),
        Inches(1.42),
        Inches(4.05),
        Inches(5.25),
        "Real merge and conflict evidence",
        [
            "Your public commit history shows merge pull requests.",
            "It also shows a real `Resolved merge conflict` commit entry.",
            "Manual conflict fixing is still done in code before commit.",
            "This makes the conflict slide based on your real GitHub activity.",
        ],
    )
    add_picture(slide, REAL / "21RABIN_billing_commits.png", Inches(4.82), Inches(1.32), width=Inches(7.3))
    add_picture(slide, GENERIC / "git_manual_conflict_terminal.png", Inches(5.8), Inches(4.8), width=Inches(6.1))

    # Slide 6
    slide = prs.slides[5]
    clear_non_title_shapes(slide)
    add_textbox(
        slide,
        Inches(0.5),
        Inches(1.45),
        Inches(4.1),
        Inches(5.2),
        "Flyway + Spring Boot together",
        [
            "Spring Boot starts the application.",
            "Flyway checks pending SQL migrations automatically.",
            "Database changes run in version order during startup.",
            "Application code and schema stay synchronized.",
        ],
    )
    add_picture(slide, GENERIC / "spring_quickstart.png", Inches(4.75), Inches(1.35), width=Inches(6.9))
    add_picture(slide, GENERIC / "spring_flyway_terminal.png", Inches(5.0), Inches(4.75), width=Inches(6.9))

    # Slide 7
    slide = prs.slides[6]
    clear_non_title_shapes(slide)
    add_textbox(
        slide,
        Inches(0.5),
        Inches(1.42),
        Inches(4.05),
        Inches(5.25),
        "Why Flyway is useful",
        [
            "Keeps schema versions in source control.",
            "Shows which migrations are pending or already applied.",
            "Reduces manual database update mistakes.",
            "Makes deployments repeatable across environments.",
        ],
    )
    add_picture(slide, GENERIC / "flyway_pending.png", Inches(5.0), Inches(1.6), width=Inches(2.8))
    add_picture(slide, GENERIC / "flyway_history.png", Inches(8.05), Inches(1.75), width=Inches(4.2))
    add_picture(slide, GENERIC / "spring_flyway_terminal.png", Inches(5.0), Inches(4.4), width=Inches(7.0))

    # Slide 8
    slide = prs.slides[7]
    clear_non_title_shapes(slide)
    add_textbox(
        slide,
        Inches(0.5),
        Inches(1.42),
        Inches(4.1),
        Inches(5.25),
        "Migration workflow example",
        [
            "Requirement arrives for a schema change.",
            "Create `V1`, `V2`, `V3` SQL files in migrations.",
            "Commit and push the code to GitHub.",
            "During deploy, Flyway executes only pending migrations.",
        ],
    )
    add_picture(slide, GENERIC / "spring_quickstart.png", Inches(4.8), Inches(1.35), width=Inches(5.8))
    add_picture(slide, GENERIC / "flyway_pending.png", Inches(10.8), Inches(1.7), width=Inches(2.0))
    add_picture(slide, GENERIC / "spring_flyway_terminal.png", Inches(4.95), Inches(4.65), width=Inches(7.0))

    # Slide 9
    slide = prs.slides[8]
    clear_non_title_shapes(slide)
    add_textbox(
        slide,
        Inches(0.5),
        Inches(1.42),
        Inches(4.1),
        Inches(5.25),
        "Real migration naming sample",
        [
            "`V1__create_users.sql`",
            "`V2__add_email.sql`",
            "`V3__create_orders.sql`",
            "Flyway records execution history in `flyway_schema_history`.",
        ],
    )
    add_picture(slide, GENERIC / "flyway_history.png", Inches(4.95), Inches(1.85), width=Inches(6.8))
    add_picture(slide, GENERIC / "flyway_pending.png", Inches(9.9), Inches(4.6), width=Inches(2.5))

    # Slide 10
    slide = prs.slides[9]
    clear_non_title_shapes(slide)
    add_textbox(
        slide,
        Inches(0.5),
        Inches(1.42),
        Inches(4.1),
        Inches(5.25),
        "Merge to main on your repo",
        [
            "Your repo page shows merged work on `main`.",
            "The commit area shows merge-related history clearly.",
            "This demonstrates code review and merge flow on GitHub.",
            "It is a real repository page from `21RABIN/billing`.",
        ],
    )
    add_picture(slide, REAL / "21RABIN_billing_repo.png", Inches(4.85), Inches(1.35), width=Inches(7.2))

    # Slide 11
    slide = prs.slides[10]
    clear_non_title_shapes(slide)
    add_textbox(
        slide,
        Inches(0.5),
        Inches(1.42),
        Inches(4.1),
        Inches(5.25),
        "Best practice reminders",
        [
            "Use one branch per feature or bug fix.",
            "Write meaningful commit messages before push.",
            "Review merge commits and conflict resolutions carefully.",
            "Never edit old Flyway migrations already run in production.",
        ],
    )
    add_picture(slide, REAL / "21RABIN_billing_commits.png", Inches(4.85), Inches(1.45), width=Inches(7.2))

    prs.save(str(OUTPUT_PPT))
    print(f"saved {OUTPUT_PPT}")


if __name__ == "__main__":
    main()
