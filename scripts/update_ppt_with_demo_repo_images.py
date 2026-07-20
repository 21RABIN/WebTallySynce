from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path("/home/user/Videos/TALLY/Tally_BackUp/Correct Working Tally Synce/tallyconnector_Jun13/tallyconnector")
SOURCE_PPT = Path("/home/user/Downloads/Source_Code_and_Database_Version_Management_Presentation.pptx")
OUTPUT_PPT = Path("/home/user/Downloads/Source_Code_and_Database_Version_Management_Presentation_demo_repo_realtime.pptx")
GENERIC = ROOT / "output" / "ppt_assets"
DEMO = ROOT / "output" / "github_demo"
REAL = ROOT / "output" / "github_real"


def clear_non_title_shapes(slide):
    title = slide.shapes.title
    for shape in list(slide.shapes):
        if shape == title:
            continue
        shape._element.getparent().remove(shape._element)


def add_textbox(slide, left, top, width, height, title, lines, title_size=26, body_size=18):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(title_size)
    p.font.bold = True
    p.font.color.rgb = RGBColor(31, 41, 55)
    p.alignment = PP_ALIGN.LEFT

    for line in lines:
        p = tf.add_paragraph()
        p.text = line
        p.level = 0
        p.font.size = Pt(body_size)
        p.font.color.rgb = RGBColor(55, 65, 81)
        p.space_after = Pt(8)


def add_picture(slide, path, left, top, width=None, height=None):
    slide.shapes.add_picture(str(path), left, top, width=width, height=height)


def main():
    prs = Presentation(str(SOURCE_PPT))

    slide = prs.slides[1]
    clear_non_title_shapes(slide)
    add_textbox(
        slide,
        Inches(0.55),
        Inches(1.45),
        Inches(4.15),
        Inches(5.3),
        "Real GitHub account view",
        [
            "Public profile used: `github.com/21RABIN`.",
            "This shows repositories and contribution activity.",
            "The presentation now uses your real GitHub account.",
            "From this account, a new demo repo was created for workflow images.",
        ],
    )
    add_picture(slide, REAL / "21RABIN_profile.png", Inches(4.85), Inches(1.4), width=Inches(7.35))

    slide = prs.slides[2]
    clear_non_title_shapes(slide)
    add_textbox(
        slide,
        Inches(0.5),
        Inches(1.42),
        Inches(4.05),
        Inches(5.25),
        "Real repo created for demo",
        [
            "New repo: `21RABIN/git-flyway-springboot-demo`.",
            "Contains Spring Boot app files and Flyway migrations.",
            "Latest commit is visible directly on the repo home page.",
            "This is a real repository created from your account access.",
        ],
    )
    add_picture(slide, DEMO / "demo_repo.png", Inches(4.85), Inches(1.35), width=Inches(7.2))

    slide = prs.slides[3]
    clear_non_title_shapes(slide)
    add_textbox(
        slide,
        Inches(0.5),
        Inches(1.42),
        Inches(4.05),
        Inches(5.25),
        "Real branch workflow",
        [
            "Branches created in the demo repo are visible here.",
            "`feature/add-email-migration` was pushed and merged.",
            "`feature/create-orders-migration` is still open for comparison.",
            "This shows the real branch view from GitHub.",
        ],
    )
    add_picture(slide, DEMO / "demo_branches.png", Inches(4.9), Inches(1.45), width=Inches(7.0))

    slide = prs.slides[4]
    clear_non_title_shapes(slide)
    add_textbox(
        slide,
        Inches(0.5),
        Inches(1.42),
        Inches(4.1),
        Inches(5.25),
        "Real commit and conflict history",
        [
            "This commit list includes the real conflict-resolution commit.",
            "`Resolve merge conflict in README` was created during setup.",
            "It also shows migration commits on `main`.",
            "This is live history from the demo repository.",
        ],
    )
    add_picture(slide, DEMO / "demo_commits.png", Inches(4.85), Inches(1.35), width=Inches(7.2))

    slide = prs.slides[5]
    clear_non_title_shapes(slide)
    add_textbox(
        slide,
        Inches(0.5),
        Inches(1.45),
        Inches(4.1),
        Inches(5.2),
        "Flyway + Spring Boot together",
        [
            "Spring Boot starts the application.",
            "Flyway checks pending SQL migrations automatically.",
            "Database changes run in version order during startup.",
            "Application code and schema stay synchronized.",
        ],
    )
    add_picture(slide, GENERIC / "spring_quickstart.png", Inches(4.75), Inches(1.35), width=Inches(6.9))
    add_picture(slide, GENERIC / "spring_flyway_terminal.png", Inches(5.0), Inches(4.75), width=Inches(6.9))

    slide = prs.slides[6]
    clear_non_title_shapes(slide)
    add_textbox(
        slide,
        Inches(0.5),
        Inches(1.42),
        Inches(4.05),
        Inches(5.25),
        "Why Flyway is useful",
        [
            "The demo repo uses versioned SQL migration files.",
            "`V1`, `V2`, and `V3` examples explain schema evolution.",
            "Flyway applies only pending migrations in order.",
            "This helps teams keep code and database changes aligned.",
        ],
    )
    add_picture(slide, GENERIC / "flyway_pending.png", Inches(5.0), Inches(1.6), width=Inches(2.8))
    add_picture(slide, GENERIC / "flyway_history.png", Inches(8.05), Inches(1.75), width=Inches(4.2))
    add_picture(slide, GENERIC / "spring_flyway_terminal.png", Inches(5.0), Inches(4.4), width=Inches(7.0))

    slide = prs.slides[7]
    clear_non_title_shapes(slide)
    add_textbox(
        slide,
        Inches(0.5),
        Inches(1.42),
        Inches(4.05),
        Inches(5.25),
        "Real pull request comparison",
        [
            "The compare page shows branch-to-main changes on GitHub.",
            "`feature/create-orders-migration` is compared with `main`.",
            "GitHub confirms the branch is able to merge.",
            "This is the closest real-time PR screen without opening a formal PR.",
        ],
    )
    add_picture(slide, DEMO / "demo_compare_pr.png", Inches(4.85), Inches(1.35), width=Inches(7.2))

    slide = prs.slides[8]
    clear_non_title_shapes(slide)
    add_textbox(
        slide,
        Inches(0.5),
        Inches(1.42),
        Inches(4.1),
        Inches(5.25),
        "Real migration naming sample",
        [
            "`V1__create_users.sql`",
            "`V2__add_email.sql`",
            "`V3__create_orders.sql`",
            "These versions are represented in the demo repo workflow.",
        ],
    )
    add_picture(slide, DEMO / "demo_compare_pr.png", Inches(4.95), Inches(1.85), width=Inches(6.8))

    slide = prs.slides[9]
    clear_non_title_shapes(slide)
    add_textbox(
        slide,
        Inches(0.5),
        Inches(1.42),
        Inches(4.05),
        Inches(5.25),
        "Code + database release flow",
        [
            "Code was committed to a real repo from your account.",
            "Feature branches were pushed to GitHub.",
            "One branch was merged to `main`; another remains for comparison.",
            "This gives a full real Git + Flyway release example.",
        ],
    )
    add_picture(slide, DEMO / "demo_repo.png", Inches(4.85), Inches(1.35), width=Inches(7.2))

    slide = prs.slides[10]
    clear_non_title_shapes(slide)
    add_textbox(
        slide,
        Inches(0.5),
        Inches(1.42),
        Inches(4.05),
        Inches(5.25),
        "Best practice reminders",
        [
            "Use one branch per feature or migration.",
            "Commit small meaningful changes.",
            "Resolve conflicts before pushing final merges.",
            "Keep Flyway versions incremental and readable.",
        ],
    )
    add_picture(slide, DEMO / "demo_branches.png", Inches(4.85), Inches(1.45), width=Inches(7.2))

    prs.save(str(OUTPUT_PPT))
    print(f"saved {OUTPUT_PPT}")


if __name__ == "__main__":
    main()
