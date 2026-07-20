from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path("/home/user/Videos/TALLY/Tally_BackUp/Correct Working Tally Synce/tallyconnector_Jun13/tallyconnector")
SOURCE_PPT = Path("/home/user/Downloads/Source_Code_and_Database_Version_Management_Presentation.pptx")
OUTPUT_PPT = Path("/home/user/Downloads/Source_Code_and_Database_Version_Management_Presentation_with_realtime_images.pptx")
ASSETS = ROOT / "output" / "ppt_assets"


def clear_non_title_shapes(slide):
    title = slide.shapes.title
    for shape in list(slide.shapes):
        if shape == title:
            continue
        shape._element.getparent().remove(shape._element)


def add_textbox(slide, left, top, width, height, title, lines, title_size=26, body_size=19):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(title_size)
    p.font.bold = True
    p.font.color.rgb = RGBColor(31, 41, 55)
    p.alignment = PP_ALIGN.LEFT

    for line in lines:
        p = tf.add_paragraph()
        p.text = line
        p.level = 0
        p.font.size = Pt(body_size)
        p.font.color.rgb = RGBColor(55, 65, 81)
        p.space_after = Pt(8)

    return box


def add_image(slide, name, left, top, width=None, height=None):
    path = ASSETS / name
    slide.shapes.add_picture(str(path), left, top, width=width, height=height)


def main():
    prs = Presentation(str(SOURCE_PPT))

    # Slide 2
    slide = prs.slides[1]
    clear_non_title_shapes(slide)
    add_textbox(
        slide,
        Inches(0.55),
        Inches(1.55),
        Inches(4.2),
        Inches(4.9),
        "Why teams use Git and GitHub",
        [
            "Track every code change with history.",
            "Commit work safely before sharing it.",
            "Push to GitHub for backup and collaboration.",
            "Review work through branches and pull requests.",
        ],
    )
    add_image(slide, "git_commit_push_terminal.png", Inches(5.05), Inches(1.45), width=Inches(7.5))

    # Slide 3
    slide = prs.slides[2]
    clear_non_title_shapes(slide)
    add_textbox(
        slide,
        Inches(0.5),
        Inches(1.45),
        Inches(4.0),
        Inches(5.2),
        "Daily workflow with real commands",
        [
            "1. `git clone` the repository",
            "2. `git checkout -b feature/ledger-sync`",
            "3. `git add` and `git commit -m \"message\"`",
            "4. `git push origin feature/ledger-sync`",
            "5. Open Pull Request on GitHub",
            "6. Review and merge to `main`",
        ],
        body_size=17,
    )
    add_image(slide, "git_commit_push_terminal.png", Inches(4.75), Inches(1.35), width=Inches(7.7))
    add_image(slide, "github_compare_pr.png", Inches(4.9), Inches(5.0), width=Inches(7.0))

    # Slide 4
    slide = prs.slides[3]
    clear_non_title_shapes(slide)
    add_textbox(
        slide,
        Inches(0.5),
        Inches(1.45),
        Inches(4.1),
        Inches(5.2),
        "Branching example",
        [
            "`main` stays stable and deployable.",
            "Each developer creates a separate feature branch.",
            "GitHub shows branches clearly before opening a PR.",
            "After review, the feature branch is merged back to `main`.",
        ],
        body_size=18,
    )
    add_image(slide, "git_branch_graph_terminal.png", Inches(4.75), Inches(1.35), width=Inches(7.6))
    add_image(slide, "github_branch_dropdown.png", Inches(8.6), Inches(4.75), width=Inches(3.4))

    # Slide 5
    slide = prs.slides[4]
    clear_non_title_shapes(slide)
    add_textbox(
        slide,
        Inches(0.5),
        Inches(1.4),
        Inches(4.1),
        Inches(5.35),
        "Manual conflict resolution",
        [
            "Conflicts happen when two branches change the same lines.",
            "GitHub warns the team before merge.",
            "Developers compare both versions and keep the correct code.",
            "After fixing, run `git add`, `git commit`, then push again.",
        ],
        body_size=18,
    )
    add_image(slide, "github_conflict_button.png", Inches(4.85), Inches(1.35), width=Inches(7.3))
    add_image(slide, "git_manual_conflict_terminal.png", Inches(4.85), Inches(2.35), width=Inches(7.3))
    add_image(slide, "github_conflict_editor.png", Inches(6.1), Inches(5.25), width=Inches(4.8))

    # Slide 6
    slide = prs.slides[5]
    clear_non_title_shapes(slide)
    add_textbox(
        slide,
        Inches(0.52),
        Inches(1.45),
        Inches(4.05),
        Inches(5.2),
        "Flyway + Spring Boot together",
        [
            "Spring Boot starts the application.",
            "Flyway checks pending SQL migrations automatically.",
            "Database changes run in version order during startup.",
            "Application code and schema stay synchronized.",
        ],
        body_size=18,
    )
    add_image(slide, "spring_quickstart.png", Inches(4.8), Inches(1.35), width=Inches(6.9))
    add_image(slide, "spring_flyway_terminal.png", Inches(5.0), Inches(4.75), width=Inches(6.9))

    # Slide 7
    slide = prs.slides[6]
    clear_non_title_shapes(slide)
    add_textbox(
        slide,
        Inches(0.52),
        Inches(1.45),
        Inches(4.05),
        Inches(5.2),
        "Why Flyway is useful",
        [
            "Keeps schema versions in source control.",
            "Shows which migrations are pending or already applied.",
            "Reduces manual database update mistakes.",
            "Makes deployments repeatable across environments.",
        ],
        body_size=18,
    )
    add_image(slide, "flyway_pending.png", Inches(5.1), Inches(1.55), width=Inches(2.9))
    add_image(slide, "flyway_history.png", Inches(8.25), Inches(1.7), width=Inches(4.1))
    add_image(slide, "spring_flyway_terminal.png", Inches(5.0), Inches(4.35), width=Inches(7.0))

    # Slide 8
    slide = prs.slides[7]
    clear_non_title_shapes(slide)
    add_textbox(
        slide,
        Inches(0.5),
        Inches(1.45),
        Inches(4.15),
        Inches(5.2),
        "Migration workflow example",
        [
            "Requirement arrives for a schema change.",
            "Create `V1`, `V2`, `V3` SQL files in the migrations folder.",
            "Commit and push the code to GitHub.",
            "During deploy, Flyway executes only pending migrations.",
        ],
        body_size=18,
    )
    add_image(slide, "spring_quickstart.png", Inches(4.8), Inches(1.35), width=Inches(5.8))
    add_image(slide, "flyway_pending.png", Inches(10.8), Inches(1.7), width=Inches(2.0))
    add_image(slide, "spring_flyway_terminal.png", Inches(4.95), Inches(4.65), width=Inches(7.0))

    # Slide 9
    slide = prs.slides[8]
    clear_non_title_shapes(slide)
    add_textbox(
        slide,
        Inches(0.5),
        Inches(1.45),
        Inches(4.1),
        Inches(5.2),
        "Real migration naming sample",
        [
            "`V1__create_users.sql`",
            "`V2__add_email.sql`",
            "`V3__create_orders.sql`",
            "Flyway records execution history in `flyway_schema_history`.",
        ],
        body_size=19,
    )
    add_image(slide, "flyway_history.png", Inches(4.95), Inches(1.85), width=Inches(6.8))
    add_image(slide, "flyway_pending.png", Inches(9.9), Inches(4.6), width=Inches(2.5))

    # Slide 10
    slide = prs.slides[9]
    clear_non_title_shapes(slide)
    add_textbox(
        slide,
        Inches(0.5),
        Inches(1.42),
        Inches(4.1),
        Inches(5.25),
        "Code + database release flow",
        [
            "Developer commits application code and migration scripts.",
            "Feature branch is pushed and reviewed on GitHub.",
            "After approval, merge to `main`.",
            "Deploy runs Spring Boot and Flyway together.",
        ],
        body_size=18,
    )
    add_image(slide, "git_pull_merge_terminal.png", Inches(4.85), Inches(1.35), width=Inches(7.2))
    add_image(slide, "github_merge_options.png", Inches(8.8), Inches(4.45), width=Inches(3.0))

    # Slide 11
    slide = prs.slides[10]
    clear_non_title_shapes(slide)
    add_textbox(
        slide,
        Inches(0.5),
        Inches(1.42),
        Inches(4.15),
        Inches(5.3),
        "Best practice reminders",
        [
            "Use one branch per feature or bug fix.",
            "Write meaningful commit messages before push.",
            "Pull latest `main` before merging large changes.",
            "Never edit old Flyway migrations already run in production.",
        ],
        body_size=18,
    )
    add_image(slide, "github_sync_fork.png", Inches(4.95), Inches(1.45), width=Inches(7.0))
    add_image(slide, "github_merge_options.png", Inches(8.9), Inches(4.45), width=Inches(2.9))

    prs.save(str(OUTPUT_PPT))
    print(f"saved {OUTPUT_PPT}")


if __name__ == "__main__":
    main()
