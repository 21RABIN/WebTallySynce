from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path("/home/user/Videos/TALLY/Tally_BackUp/Correct Working Tally Synce/tallyconnector_Jun13/tallyconnector")
SOURCE_PPT = Path("/home/user/Downloads/Source_Code_and_Database_Version_Management_Presentation.pptx")
OUTPUT_PPT = Path("/home/user/Downloads/Source_Code_and_Database_Version_Management_Presentation_full_realtime_images.pptx")
DEMO = ROOT / "output" / "github_demo"
GENERIC = ROOT / "output" / "ppt_assets"
REAL = ROOT / "output" / "github_real"


def clear_non_title_shapes(slide):
    title = slide.shapes.title
    for shape in list(slide.shapes):
        if shape == title:
            continue
        shape._element.getparent().remove(shape._element)


def add_textbox(slide, left, top, width, height, title, lines, title_size=25, body_size=18):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(title_size)
    p.font.bold = True
    p.font.color.rgb = RGBColor(31, 41, 55)
    p.alignment = PP_ALIGN.LEFT

    for line in lines:
        p = tf.add_paragraph()
        p.text = line
        p.level = 0
        p.font.size = Pt(body_size)
        p.font.color.rgb = RGBColor(55, 65, 81)
        p.space_after = Pt(6)


def add_picture(slide, path, left, top, width=None, height=None):
    slide.shapes.add_picture(str(path), left, top, width=width, height=height)


def main():
    prs = Presentation(str(SOURCE_PPT))

    slide = prs.slides[1]
    clear_non_title_shapes(slide)
    add_textbox(
        slide, Inches(0.5), Inches(1.45), Inches(4.1), Inches(5.2),
        "Real GitHub account and repo",
        [
            "Account used: `github.com/21RABIN`.",
            "Real demo repo was created from this account.",
            "Teams can see repositories, activity, and ownership.",
            "This makes the presentation use live GitHub examples.",
        ],
    )
    add_picture(slide, REAL / "21RABIN_profile.png", Inches(4.8), Inches(1.4), width=Inches(7.25))

    slide = prs.slides[2]
    clear_non_title_shapes(slide)
    add_textbox(
        slide, Inches(0.45), Inches(1.4), Inches(3.9), Inches(5.3),
        "How to commit and push code",
        [
            "Create a feature branch.",
            "Add the changed file.",
            "Commit with a meaningful message.",
            "Push the branch to GitHub.",
            "GitHub gives the pull request link.",
        ],
    )
    add_picture(slide, DEMO / "real_commit_push_terminal.png", Inches(4.35), Inches(1.35), width=Inches(7.6))

    slide = prs.slides[3]
    clear_non_title_shapes(slide)
    add_textbox(
        slide, Inches(0.45), Inches(1.4), Inches(3.9), Inches(5.3),
        "Real branches view",
        [
            "This page shows `main` and feature branches.",
            "One branch was merged.",
            "One branch is still open for pull request work.",
            "This is the real branch page from GitHub.",
        ],
    )
    add_picture(slide, DEMO / "demo_branches.png", Inches(4.45), Inches(1.45), width=Inches(7.35))

    slide = prs.slides[4]
    clear_non_title_shapes(slide)
    add_textbox(
        slide, Inches(0.45), Inches(1.4), Inches(3.95), Inches(5.3),
        "Manual conflict example",
        [
            "A real merge conflict happened in `README.md`.",
            "Git history shows the conflict-resolution commit.",
            "Manual conflict fixing is done in the code, then committed.",
            "This gives both real GitHub history and real terminal flow.",
        ],
    )
    add_picture(slide, DEMO / "demo_commits.png", Inches(4.55), Inches(1.35), width=Inches(7.2))
    add_picture(slide, DEMO / "real_manual_conflict_terminal.png", Inches(5.2), Inches(4.45), width=Inches(6.6))

    slide = prs.slides[5]
    clear_non_title_shapes(slide)
    add_textbox(
        slide, Inches(0.45), Inches(1.45), Inches(3.95), Inches(5.2),
        "Flyway and Spring Boot",
        [
            "Spring Boot starts the application.",
            "Flyway automatically checks pending SQL migrations.",
            "Database schema changes run in version order.",
            "Code and database stay synchronized.",
        ],
    )
    add_picture(slide, GENERIC / "spring_quickstart.png", Inches(4.55), Inches(1.35), width=Inches(6.8))
    add_picture(slide, GENERIC / "spring_flyway_terminal.png", Inches(4.95), Inches(4.8), width=Inches(7.0))

    slide = prs.slides[6]
    clear_non_title_shapes(slide)
    add_textbox(
        slide, Inches(0.45), Inches(1.4), Inches(3.95), Inches(5.25),
        "Why Flyway is useful",
        [
            "Migration files are stored in Git.",
            "Versions like `V1`, `V2`, `V3` are easy to track.",
            "Only pending migrations are applied.",
            "This helps safe deployment in all environments.",
        ],
    )
    add_picture(slide, GENERIC / "flyway_pending.png", Inches(4.9), Inches(1.6), width=Inches(2.8))
    add_picture(slide, GENERIC / "flyway_history.png", Inches(7.9), Inches(1.75), width=Inches(4.25))
    add_picture(slide, GENERIC / "spring_flyway_terminal.png", Inches(4.95), Inches(4.4), width=Inches(7.0))

    slide = prs.slides[7]
    clear_non_title_shapes(slide)
    add_textbox(
        slide, Inches(0.45), Inches(1.35), Inches(3.85), Inches(5.4),
        "Real pull request creation",
        [
            "This is the actual GitHub pull request creation screen.",
            "Branch `feature/create-orders-migration` is compared with `main`.",
            "GitHub shows that the branch is able to merge.",
            "This matches the real workflow after `git push`.",
        ],
    )
    add_picture(slide, DEMO / "demo_pr_create.png", Inches(4.35), Inches(1.25), width=Inches(7.7))

    slide = prs.slides[8]
    clear_non_title_shapes(slide)
    add_textbox(
        slide, Inches(0.45), Inches(1.4), Inches(3.95), Inches(5.3),
        "Real migration file example",
        [
            "`V1__create_users.sql`",
            "`V2__add_email.sql`",
            "`V3__create_orders.sql`",
            "The compare page shows the actual SQL file content added in GitHub.",
        ],
    )
    add_picture(slide, DEMO / "demo_compare_pr.png", Inches(4.55), Inches(1.35), width=Inches(7.25))

    slide = prs.slides[9]
    clear_non_title_shapes(slide)
    add_textbox(
        slide, Inches(0.45), Inches(1.4), Inches(3.95), Inches(5.3),
        "How to pull and merge to main",
        [
            "Checkout `main`.",
            "Pull the latest code from GitHub.",
            "Merge the feature branch.",
            "Push `main` back to GitHub.",
            "This is the real terminal workflow for merge to main.",
        ],
    )
    add_picture(slide, DEMO / "real_pull_merge_terminal.png", Inches(4.45), Inches(1.45), width=Inches(7.45))

    slide = prs.slides[10]
    clear_non_title_shapes(slide)
    add_textbox(
        slide, Inches(0.45), Inches(1.4), Inches(3.95), Inches(5.3),
        "Best practice reminders",
        [
            "One branch for one feature or migration.",
            "Use clear commit messages.",
            "Review PR before merge.",
            "Resolve conflicts carefully.",
            "Keep Flyway version numbers clean and ordered.",
        ],
    )
    add_picture(slide, DEMO / "demo_repo.png", Inches(4.45), Inches(1.35), width=Inches(7.35))

    prs.save(str(OUTPUT_PPT))
    print(f"saved {OUTPUT_PPT}")


if __name__ == "__main__":
    main()
